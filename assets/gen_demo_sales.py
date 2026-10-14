# -*- coding: utf-8 -*-
"""生成 README 首屏演示数据表 assets/demo-sales.xlsx。

视频演示链路的道具文件(Excel 左半屏 → `actl extract role:DataGrid` →
JSON → 汇总句写进 Word),三条硬约束决定了表的结构:

1. 表格从 A1 起、无标题行、无合计行、无留白边距 —— extract 读的是整个
   used range(docs/14:全网格输出,空单元格回填坐标名),标题/边距会以
   坐标名噪音混进 JSON。屏幕所见即 JSON 所得,used range 必须恰为 A1:D8。
2. 不写合计行 —— "10月合计 84,000"由 agent 从 7 行数据算出,这是演示
   叙事的一部分(agent 真的对数据做了加工),不能让表替它算好。
3. 数值全静态、零公式 —— 冻结的演示道具,不做 Live Formula;也因此
   无需 LibreOffice recalc(audit/scan/validate 照跑)。

数据设计成可当场对账:9 月合计 75,000,10 月合计 84,000,
整体环比 = 9,000/75,000 = +12.0% 整。视频汇总文案即取此:
「10月合计营收 84,000 元,环比 +12%,七大区域全部增长。」

依赖:xlsx 技能环境(templates/base.py 设计令牌)+ openpyxl。
"""

import os
import sys

XLSX_SKILL_DIR = os.environ.get(
    "XLSX_SKILL_DIR",
    r"C:\Users\gqy17\.zcode\cli\plugins\cache\zcode-plugins-official"
    r"\spreadsheets\0.1.8\skills\xlsx",
)
for sub in (XLSX_SKILL_DIR, os.path.join(XLSX_SKILL_DIR, "templates")):
    if sub not in sys.path:
        sys.path.insert(0, sub)

import base
from base import FORMATS

from openpyxl import Workbook, load_workbook

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "demo-sales.xlsx")

# 本地字体优先(SKILL 环境原则):zh-CN Windows 必装微软雅黑,未装则回落
# base.py 默认解析链(Noto Sans SC,Excel 渲染时自行替换)。
if os.path.exists(os.path.join(os.environ.get("WINDIR", r"C:\Windows"), "Fonts", "msyh.ttc")):
    base.FONT_NAME = "Microsoft YaHei"
    base.HEADER_BOLD = True

# (区域, 9月营收, 10月营收) —— 列合计 75,000 / 84,000,环比恰为 +12.0%
REGIONS = [
    ("华东", 19800, 21500),
    ("华北", 15500, 16800),
    ("华南", 13400, 14600),
    ("华中", 9600, 11200),
    ("西南", 8200, 9300),
    ("西北", 5300, 6100),
    ("东北", 3200, 4500),
]
HEADERS = ["区域", "9月营收（元）", "10月营收（元）", "环比增幅"]

wb = Workbook()
ws = wb.active
ws.title = "10月销售汇总"
ws.sheet_view.showGridLines = False

for col, header in enumerate(HEADERS, 1):
    ws.cell(row=1, column=col, value=header)
base.style_header_row(ws, row_num=1, col_start=1, col_end=4)

for i, (region, sep_rev, oct_rev) in enumerate(REGIONS):
    row = 2 + i
    ws.cell(row=row, column=1, value=region)
    ws.cell(row=row, column=2, value=sep_rev)
    ws.cell(row=row, column=3, value=oct_rev)
    ws.cell(row=row, column=4, value=(oct_rev - sep_rev) / sep_rev)
    base.style_data_row(ws, row_num=row, col_start=1, col_end=4, row_index=i)
    for col in (2, 3, 4):
        cell = ws.cell(row=row, column=col)
        cell.number_format = FORMATS["integer"] if col < 4 else FORMATS["percentage"]
        cell.alignment = base.align_number()

base.auto_fit_columns(ws, min_width=8, max_width=28, header_row=1, data_start_row=2)
# auto_fit 按数据估宽,中文表头更宽;演示画面要求表头单行不折,手动放宽
for col_letter, width in (("A", 10), ("B", 14), ("C", 14), ("D", 11)):
    ws.column_dimensions[col_letter].width = width
base.auto_fit_row_heights(ws, header_row=1, data_start_row=2)

wb.properties.creator = "Z.ai"
wb.save(OUT)

# ---- Inspector:读回语义核验(独立于写入路径) ----
wb2 = load_workbook(OUT)
ws2 = wb2.active
assert ws2.dimensions == "A1:D8", f"used range 应恰为 A1:D8,实际 {ws2.dimensions}"
assert [c.value for c in ws2[1]] == HEADERS, "表头读回不一致"

sep_total = sum(r[1] for r in REGIONS)
oct_total = sum(r[2] for r in REGIONS)
read_sep = sum(ws2.cell(row=r, column=2).value for r in range(2, 9))
read_oct = sum(ws2.cell(row=r, column=3).value for r in range(2, 9))
assert read_sep == sep_total == 75000, f"9月合计 {read_sep} != {sep_total}"
assert read_oct == oct_total == 84000, f"10月合计 {read_oct} != {oct_total}"
growth_total = (oct_total - sep_total) / sep_total
assert growth_total == 0.12, f"整体环比应为 +12.0%,实际 {growth_total}"

for i, (region, sep_rev, oct_rev) in enumerate(REGIONS):
    row = 2 + i
    assert ws2.cell(row=row, column=1).value == region
    assert ws2.cell(row=row, column=2).value == sep_rev
    assert ws2.cell(row=row, column=3).value == oct_rev
    assert abs(ws2.cell(row=row, column=4).value - (oct_rev - sep_rev) / sep_rev) < 1e-12

print(f"OK: {OUT}")
print(f"used range: {ws2.dimensions}, sheet: {ws2.title}")
print(f"9月合计 {read_sep:,} / 10月合计 {read_oct:,} / 整体环比 +{growth_total:.1%}")
print("视频汇总文案: 10月合计营收 84,000 元,环比 +12%,七大区域全部增长。")
